"""tests/integration/test_workspace_api.py — workspace endpoints.

Requires a running harness agent (``env_with_agent``); workspace I/O goes
through ``agent.workspace`` backed by ``local_shell`` on the agent dir.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from docx import Document

# Workspace UI semantics: leading '/' is relative to agent workspace.
FROM_WORKSPACE = {"from_workspace": "true"}


def _sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Report Title", level=1)
    doc.add_paragraph("Intro paragraph")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
async def env(env_with_agent):
    yield env_with_agent


# --- listing ---------------------------------------------------------------


async def test_tree_returns_empty_for_fresh_workspace(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(f"/api/agents/{aid}/workspace/tree?from_workspace=true", headers=auth)
    assert r.status_code == 200, r.text
    rows = r.json()
    assert isinstance(rows, list)
    # Fresh workspace may contain a SOUL.md (written at agent boot) or
    # be empty — we don't pin the exact contents, just the shape.
    for row in rows:
        assert "path" in row


async def test_tree_lists_root_files(env: Any) -> None:
    c, srv, auth, aid = env
    agent = srv.app_runtime.agent_registry.get_agent(aid)
    await agent.workspace.aupload_bytes("notes.md", b"hello")

    r = await c.get(f"/api/agents/{aid}/workspace/tree?path=/&from_workspace=true", headers=auth)
    assert r.status_code == 200, r.text
    rows = r.json()
    paths = {row["path"] for row in rows}
    assert any("notes.md" in p for p in paths)

    r = await c.get(
        f"/api/agents/{aid}/workspace/file?path=%2Fnotes.md&from_workspace=true",
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["content"] == "hello"


async def test_tree_lists_subdirectory(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/sub/nested.txt"},
        headers=auth,
        json={"content": "nested content"},
    )

    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/sub"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    rows = r.json()
    assert any("nested.txt" in row["path"] for row in rows)

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/sub/nested.txt"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["content"] == "nested content"


async def test_tree_for_unknown_agent_404(env: Any) -> None:
    c, _srv, auth, _aid = env
    r = await c.get("/api/agents/no-such-agent/workspace/tree?from_workspace=true", headers=auth)
    assert r.status_code == 404


# --- write + read round-trip ------------------------------------------------


async def test_write_then_read_roundtrip(env: Any) -> None:
    c, _srv, auth, aid = env
    payload = "hello from workspace test\n"
    r = await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/notes.md"},
        headers=auth,
        json={"content": payload},
    )
    assert r.status_code == 200, r.text
    assert r.json()["path"] == "/notes.md"
    assert r.json()["size"] == len(payload.encode("utf-8"))

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/notes.md"},
        headers=auth,
    )
    assert r.status_code == 200
    body = r.json()
    assert body["path"] == "/notes.md"
    assert body["content"] == payload


async def test_read_missing_file_404(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/no-such-file.txt"},
        headers=auth,
    )
    assert r.status_code == 404


# --- upload + download ------------------------------------------------------


async def test_upload_then_download_binary(env: Any) -> None:
    c, _srv, auth, aid = env
    blob = b"\x89PNG\r\n\x1a\n" + b"\x00" * 16  # fake PNG header
    files = {"file": ("logo.png", blob, "image/png")}
    r = await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params={**FROM_WORKSPACE},
        headers=auth,
        files=files,
    )
    assert r.status_code == 200, r.text
    assert r.json()["size"] == len(blob)

    r = await c.get(
        f"/api/agents/{aid}/workspace/download",
        params={**FROM_WORKSPACE, "path": "/logo.png"},
        headers=auth,
    )
    assert r.status_code == 200
    assert r.content.startswith(b"\x89PNG\r\n\x1a\n")
    cd = r.headers.get("content-disposition", "")
    assert "logo.png" in cd


async def test_download_non_ascii_filename(env: Any) -> None:
    c, _srv, auth, aid = env
    fname = "1783510288_地球介绍.pptx"
    path = f"/outbound/{fname}"
    r = await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params={**FROM_WORKSPACE, "path": path},
        headers=auth,
        files={"file": (fname, b"PK\x03\x04fake", "application/vnd.ms-powerpoint")},
    )
    assert r.status_code == 200, r.text

    r = await c.get(
        f"/api/agents/{aid}/workspace/download",
        params={**FROM_WORKSPACE, "path": path},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"PK\x03\x04")
    cd = r.headers.get("content-disposition", "")
    assert 'filename="download.pptx"' in cd
    assert "filename*" in cd
    assert "%E5%9C%B0%E7%90%83" in cd


async def test_upload_with_explicit_path_query(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.post(
        f"/api/agents/{aid}/workspace/upload",
        params={**FROM_WORKSPACE, "path": "/sub/dir/named.txt"},
        headers=auth,
        files={"file": ("ignored.txt", b"x", "text/plain")},
    )
    assert r.status_code == 200
    assert r.json()["path"] == "/sub/dir/named.txt"


# --- glob + grep ------------------------------------------------------------


async def test_glob_after_seeding(env: Any) -> None:
    c, _srv, auth, aid = env
    for fname in ("a.md", "b.md", "c.txt"):
        await c.put(
            f"/api/agents/{aid}/workspace/file",
            params={**FROM_WORKSPACE, "path": f"/{fname}"},
            headers=auth,
            json={"content": "x"},
        )
    r = await c.get(
        f"/api/agents/{aid}/workspace/glob",
        params={**FROM_WORKSPACE, "pattern": "*.md", "path": "/"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    paths = {row["path"] for row in r.json()}
    # Glob may return absolute or relative paths depending on backend.
    matched = {p.rsplit("/", 1)[-1] for p in paths}
    assert "a.md" in matched
    assert "b.md" in matched


async def test_grep_after_seeding(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/needle.txt"},
        headers=auth,
        json={"content": "alpha\nNEEDLE here\ngamma\n"},
    )
    r = await c.get(
        f"/api/agents/{aid}/workspace/grep",
        params={**FROM_WORKSPACE, "pattern": "NEEDLE", "path": "/"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    rows = r.json()
    assert any("NEEDLE" in str(row) for row in rows)


# --- cross-user isolation ---------------------------------------------------


async def test_non_owner_cannot_access_workspace(env: Any) -> None:
    """Non-owners cannot read another user's agent workspace."""
    c, _srv, admin_auth, _aid = env
    await c.post(
        "/api/users",
        headers=admin_auth,
        json={"username": "bob", "password": "TestPass12", "role": "user"},
    )
    bob_tok = (
        await c.post(
            "/api/auth/login",
            json={"username": "bob", "password": "TestPass12"},
        )
    ).json()["access_token"]
    bob_auth = {"Authorization": f"Bearer {bob_tok}"}

    admin_agent_id = (await c.get("/api/agents", headers=admin_auth)).json()[0]["agent_id"]

    r = await c.get(
        f"/api/agents/{admin_agent_id}/workspace/tree",
        headers=bob_auth,
    )
    assert r.status_code == 403


# --- delete + move ----------------------------------------------------------


async def test_delete_file(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/trash-me.txt"},
        headers=auth,
        json={"content": "bye"},
    )
    r = await c.delete(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/trash-me.txt"},
        headers=auth,
    )
    assert r.status_code == 204, r.text

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/trash-me.txt"},
        headers=auth,
    )
    assert r.status_code == 404


async def test_move_file(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/src.txt"},
        headers=auth,
        json={"content": "payload"},
    )
    r = await c.post(
        f"/api/agents/{aid}/workspace/move",
        params={**FROM_WORKSPACE, "path": "/src.txt"},
        headers=auth,
        json={"destination": "/moved/src.txt"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["path"] == "/moved/src.txt"

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/moved/src.txt"},
        headers=auth,
    )
    assert r.status_code == 200
    assert r.json()["content"] == "payload"

    r = await c.get(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/src.txt"},
        headers=auth,
    )
    assert r.status_code == 404


async def test_rename_file(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/old-name.md"},
        headers=auth,
        json={"content": "x"},
    )
    r = await c.post(
        f"/api/agents/{aid}/workspace/move",
        params={**FROM_WORKSPACE, "path": "/old-name.md"},
        headers=auth,
        json={"destination": "/new-name.md"},
    )
    assert r.status_code == 200, r.text
    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/"},
        headers=auth,
    )
    paths = {row["path"].rsplit("/", 1)[-1] for row in r.json()}
    assert "new-name.md" in paths
    assert "old-name.md" not in paths


async def test_mkdir_creates_directory(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.post(
        f"/api/agents/{aid}/workspace/mkdir",
        params={**FROM_WORKSPACE, "path": "/projects/demo"},
        headers=auth,
    )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["path"] == "/projects/demo"
    assert body["is_dir"] is True

    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/projects"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    names = {row["path"].rsplit("/", 1)[-1] for row in r.json()}
    assert "demo" in names


async def test_delete_directory(env: Any) -> None:
    c, _srv, auth, aid = env
    await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/box/a.txt"},
        headers=auth,
        json={"content": "a"},
    )
    r = await c.delete(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/box"},
        headers=auth,
    )
    assert r.status_code == 204, r.text
    r = await c.get(
        f"/api/agents/{aid}/workspace/tree",
        params={**FROM_WORKSPACE, "path": "/"},
        headers=auth,
    )
    names = {row["path"].rsplit("/", 1)[-1] for row in r.json()}
    assert "box" not in names


async def test_delete_builtin_skills_forbidden(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.delete(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/_builtin_skills/foo/SKILL.md"},
        headers=auth,
    )
    assert r.status_code == 403


# --- editable document (Markdown round-trip) --------------------------------


async def test_doc_read_and_write_roundtrip(env: Any) -> None:
    c, srv, auth, aid = env
    agent = srv.app_runtime.agent_registry.get_agent(aid)
    await agent.workspace.aupload_bytes("report.docx", _sample_docx_bytes())

    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/report.docx"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["path"] == "/report.docx"
    assert "# Report Title" in body["content"]
    assert "Intro paragraph" in body["content"]

    r = await c.put(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/report.docx"},
        headers=auth,
        json={"content": "# Updated Title\n\nNew **bold** body\n"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["path"] == "/report.docx"
    assert r.json()["size"] > 0

    # Stored bytes parse back into a valid docx.
    blob = await agent.workspace.adownload_bytes("report.docx")
    assert blob is not None
    parsed = Document(BytesIO(blob))
    texts = [p.text for p in parsed.paragraphs if p.text]
    assert "Updated Title" in texts
    assert "New bold body" in texts

    # Round-trip back through the API keeps the Markdown structure.
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/report.docx"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    content = r.json()["content"]
    assert "# Updated Title" in content
    assert "**bold**" in content


async def test_doc_unsupported_extension_400(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/notes.md"},
        headers=auth,
    )
    assert r.status_code == 400


async def test_create_empty_docx_via_text_endpoint_is_previewable(env: Any) -> None:
    """Workspace "new file" creates .docx through the text endpoint; it must be
    stored as a valid document package so preview/edit work immediately."""
    c, srv, auth, aid = env
    r = await c.put(
        f"/api/agents/{aid}/workspace/file",
        params={**FROM_WORKSPACE, "path": "/fresh.docx"},
        headers=auth,
        json={"content": ""},
    )
    assert r.status_code == 200, r.text
    assert r.json()["size"] > 0  # not a 0-byte file

    agent = srv.app_runtime.agent_registry.get_agent(aid)
    blob = await agent.workspace.adownload_bytes("fresh.docx")
    assert blob is not None
    parsed = Document(BytesIO(blob))
    assert len(parsed.paragraphs) == 0  # valid empty document

    # It can be opened for editing as an empty Markdown document.
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/fresh.docx"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["content"] == ""


async def test_doc_missing_file_404(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/no-such.docx"},
        headers=auth,
    )
    assert r.status_code == 404


async def test_doc_write_root_forbidden(env: Any) -> None:
    c, _srv, auth, aid = env
    r = await c.put(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/"},
        headers=auth,
        json={"content": "# hi\n"},
    )
    assert r.status_code == 403


async def test_doc_write_invalid_content_400(env: Any) -> None:
    c, srv, auth, aid = env
    agent = srv.app_runtime.agent_registry.get_agent(aid)
    await agent.workspace.aupload_bytes("broken.docx", b"not a real docx zip")

    r = await c.get(
        f"/api/agents/{aid}/workspace/doc",
        params={**FROM_WORKSPACE, "path": "/broken.docx"},
        headers=auth,
    )
    assert r.status_code == 400
