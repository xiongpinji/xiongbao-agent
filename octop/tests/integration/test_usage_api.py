"""tests/integration/test_usage_api.py — token usage ledger + summary endpoint.

Three layers:
  1. UsageRepo round-trips rows and aggregates per granularity.
  2. /api/usage/summary respects ?as_user= scope (admin) and pins
     non-admins to their own user_id.
  3. /api/admin/usage/summary returns global rollups.

The chat-stream → ledger pipe is exercised in test_chat_ws via the
existing FakeHarnessAgent fixtures; here we drive the repo directly so
the test stays focused.
"""

from __future__ import annotations

import time
from typing import Any

import pytest


@pytest.fixture
async def env(env_usage):
    _c, srv, _admin_auth, _alice_auth, ctx = env_usage
    # The usage_log table has FOREIGN KEYs to agents(agent_id) and users(id).
    # These tests log against synthetic agent ids, so seed matching parent
    # rows up front (users alice_id / 1 already exist from bootstrap).
    _seed_usage_agents(srv, ["agt1", "agt-a", "agt-b", "x", "y"], user_id=ctx["alice_id"])
    yield env_usage


def _seed_usage_agents(srv: Any, agent_ids: list[str], *, user_id: int) -> None:
    with srv.services.db.connect() as conn:
        now = int(time.time())
        for aid in agent_ids:
            conn.execute(
                "INSERT OR IGNORE INTO agents (agent_id, user_id, name, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?)",
                (aid, user_id, aid, now, now),
            )


# --- repo direct -------------------------------------------------------------


async def test_repo_record_and_summary_total(env: Any) -> None:
    _c, srv, _admin_auth, _alice_auth, ctx = env
    repo = srv.services.usage_repo
    repo.record(
        agent_id="agt1",
        user_id=ctx["alice_id"],
        thread_id="t1",
        model="openai:gpt-4o-mini",
        input_tokens=100,
        uncached_input_tokens=30,
        cache_read_tokens=70,
        output_tokens=50,
    )
    repo.record(
        agent_id="agt1",
        user_id=ctx["alice_id"],
        thread_id="t1",
        model="openai:gpt-4o-mini",
        input_tokens=200,
        output_tokens=80,
    )
    result = repo.summary(user_id=ctx["alice_id"], window="last_30d", granularity="total")
    assert result["input_tokens"] == 300
    assert result["uncached_input_tokens"] == 230
    assert result["cache_read_tokens"] == 70
    assert result["cache_hit_percent"] == 23
    assert result["output_tokens"] == 130
    assert result["total_tokens"] == 430
    assert result["turns"] == 2
    assert result["avg_per_turn"] == 215


async def test_repo_summary_by_agent(env: Any) -> None:
    _c, srv, _admin_auth, _alice_auth, ctx = env
    repo = srv.services.usage_repo
    repo.record(agent_id="agt-a", user_id=ctx["alice_id"], input_tokens=10, output_tokens=5)
    repo.record(agent_id="agt-a", user_id=ctx["alice_id"], input_tokens=20, output_tokens=10)
    repo.record(agent_id="agt-b", user_id=ctx["alice_id"], input_tokens=5, output_tokens=2)
    result = repo.summary(user_id=ctx["alice_id"], window="last_30d", granularity="by_agent")
    by_id = {b["key"]: b for b in result["buckets"]}
    assert by_id["agt-a"]["total_tokens"] == 45
    assert by_id["agt-b"]["total_tokens"] == 7


async def test_repo_summary_by_model(env: Any) -> None:
    _c, srv, _admin_auth, _alice_auth, ctx = env
    repo = srv.services.usage_repo
    repo.record(
        agent_id="x",
        user_id=ctx["alice_id"],
        model="openai:gpt-4o-mini",
        input_tokens=10,
        output_tokens=5,
    )
    repo.record(
        agent_id="x",
        user_id=ctx["alice_id"],
        model="anthropic:claude-haiku",
        input_tokens=20,
        output_tokens=10,
    )
    repo.record(
        agent_id="x",
        user_id=ctx["alice_id"],
        model="",  # unknown — should bucket as ``(unknown)``
        input_tokens=1,
        output_tokens=1,
    )
    result = repo.summary(user_id=ctx["alice_id"], window="last_30d", granularity="by_model")
    by_key = {b["key"]: b for b in result["buckets"]}
    assert by_key["openai:gpt-4o-mini"]["turns"] == 1
    assert by_key["anthropic:claude-haiku"]["total_tokens"] == 30
    assert "(unknown)" in by_key


# --- /api/usage/summary -----------------------------------------------------


async def test_user_summary_only_sees_own_data(env: Any) -> None:
    """Non-admin caller is implicitly scoped to their own user_id."""
    c, srv, admin_auth, alice_auth, ctx = env
    repo = srv.services.usage_repo
    # admin's own row (user_id=1)
    repo.record(agent_id="x", user_id=1, input_tokens=1000, output_tokens=500)
    # alice's row
    repo.record(agent_id="y", user_id=ctx["alice_id"], input_tokens=10, output_tokens=5)

    r = await c.get("/api/usage/summary?granularity=total", headers=alice_auth)
    assert r.status_code == 200
    body = r.json()
    assert body["total_tokens"] == 15
    assert body["turns"] == 1


async def test_admin_can_query_other_user_via_as_user(env: Any) -> None:
    c, srv, admin_auth, _alice_auth, ctx = env
    repo = srv.services.usage_repo
    repo.record(agent_id="y", user_id=ctx["alice_id"], input_tokens=10, output_tokens=5)

    r = await c.get(
        f"/api/usage/summary?granularity=total&as_user={ctx['alice_id']}",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 15


async def test_non_admin_as_user_is_forbidden(env: Any) -> None:
    c, _srv, _admin_auth, alice_auth, _ctx = env
    r = await c.get(
        "/api/usage/summary?as_user=1",
        headers=alice_auth,
    )
    assert r.status_code == 403


async def test_admin_summary_global(env: Any) -> None:
    c, srv, admin_auth, _alice_auth, ctx = env
    repo = srv.services.usage_repo
    repo.record(agent_id="x", user_id=1, input_tokens=100, output_tokens=50)
    repo.record(agent_id="y", user_id=ctx["alice_id"], input_tokens=10, output_tokens=5)
    r = await c.get(
        "/api/admin/usage/summary?granularity=total",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 165


async def test_admin_summary_requires_admin(env: Any) -> None:
    c, _srv, _admin_auth, alice_auth, _ctx = env
    r = await c.get("/api/admin/usage/summary", headers=alice_auth)
    assert r.status_code == 403


async def test_summary_day_window_filters_rows(env: Any) -> None:
    from datetime import datetime
    from zoneinfo import ZoneInfo

    c, srv, _admin_auth, alice_auth, ctx = env
    repo = srv.services.usage_repo
    tz = ZoneInfo("Asia/Shanghai")
    day_ts = int(datetime(2026, 9, 7, 12, 0, tzinfo=tz).timestamp())
    other_ts = int(datetime(2026, 9, 8, 12, 0, tzinfo=tz).timestamp())
    repo.record(
        agent_id="y",
        user_id=ctx["alice_id"],
        input_tokens=10,
        output_tokens=5,
        ts=day_ts,
    )
    repo.record(
        agent_id="y",
        user_id=ctx["alice_id"],
        input_tokens=100,
        output_tokens=50,
        ts=other_ts,
    )
    r = await c.get(
        "/api/usage/summary?granularity=total&window=day:2026-09-07",
        headers=alice_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 15
    assert r.json()["turns"] == 1


async def test_user_export_xlsx(env: Any) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    c, srv, _admin_auth, alice_auth, ctx = env
    repo = srv.services.usage_repo
    # Give the seeded agent a display name for the Excel column.
    with srv.services.db.connect() as conn:
        conn.execute(
            "UPDATE agents SET name = ? WHERE agent_id = ?",
            ("Export Expert", "y"),
        )
    repo.record(
        agent_id="y",
        user_id=ctx["alice_id"],
        thread_id="t-export",
        model="openai:gpt-4o-mini",
        input_tokens=11,
        output_tokens=7,
        ts=1_700_000_100,
    )
    repo.record(
        agent_id="y",
        user_id=ctx["alice_id"],
        thread_id="t-export",
        model="openai:gpt-4o-mini",
        input_tokens=3,
        output_tokens=1,
        ts=1_700_000_000,
    )
    r = await c.get(
        "/api/usage/export.xlsx?window=all",
        headers={**alice_auth, "Accept-Language": "zh"},
    )
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    disposition = r.headers.get("content-disposition", "")
    assert "filename*" in disposition
    assert "_all.xlsx" in disposition
    wb = load_workbook(BytesIO(r.content))
    assert "明细" in wb.sheetnames
    assert "按天" in wb.sheetnames
    assert "按专家" in wb.sheetnames
    assert "按模型" in wb.sheetnames
    detail = wb["明细"]
    headers = [cell.value for cell in detail[1]]
    assert "专家名称" in headers
    name_col = headers.index("专家名称") + 1
    assert any(row[name_col - 1].value == "Export Expert" for row in detail.iter_rows(min_row=2))
    time_col = headers.index(next(h for h in headers if h and h.startswith("时间"))) + 1
    # TOTAL sits two rows below the last data row (blank spacer in between)
    # so Excel Sort does not treat 合计 as part of the contiguous data block.
    assert detail.cell(detail.max_row, 1).value == "合计"
    assert all(detail.cell(detail.max_row - 1, col).value is None for col in range(1, 5))
    assert detail.auto_filter.ref == f"A1:P{detail.max_row - 2}"
    time_vals = [
        row[time_col - 1].value
        for row in detail.iter_rows(min_row=2, max_row=detail.max_row - 2)
        if row[time_col - 1].value
    ]
    assert len(time_vals) >= 2
    assert time_vals == sorted(time_vals)
    assert time_vals[0] < time_vals[-1]
    assert all(
        isinstance(v, str) and len(v) == 19 and v[4] == "-" and v[10] == " " and "+" not in v
        for v in time_vals
    )
    assert wb["按天"]._charts
    assert wb["按专家"]._charts
    assert wb["按模型"]._charts


async def test_admin_export_requires_admin(env: Any) -> None:
    c, _srv, _admin_auth, alice_auth, _ctx = env
    r = await c.get("/api/admin/usage/export.xlsx", headers=alice_auth)
    assert r.status_code == 403


async def test_invalid_day_window_returns_400(env: Any) -> None:
    c, _srv, _admin_auth, alice_auth, _ctx = env
    r = await c.get(
        "/api/usage/summary?window=day:2026-9-7",
        headers=alice_auth,
    )
    assert r.status_code == 400


async def test_admin_filters_by_user_agent_and_windows(env: Any) -> None:
    """Admin user/agent/window filters must isolate seeded multi-user rows."""
    from datetime import datetime
    from io import BytesIO
    from zoneinfo import ZoneInfo

    from openpyxl import load_workbook

    from tests.support.auth import create_user, resolve_user_id

    c, srv, admin_auth, alice_auth, ctx = env
    bob_auth = await create_user(c, admin_auth, username="usage_bob")
    bob_id = await resolve_user_id(c, admin_auth, "usage_bob")
    alice_id = ctx["alice_id"]
    _seed_usage_agents(srv, ["agt-alice", "agt-bob"], user_id=alice_id)

    repo = srv.services.usage_repo
    tz = ZoneInfo("Asia/Shanghai")
    day_a = int(datetime(2026, 8, 20, 12, 0, tzinfo=tz).timestamp())
    day_b = int(datetime(2026, 8, 25, 12, 0, tzinfo=tz).timestamp())
    day_c = int(datetime(2026, 9, 3, 12, 0, tzinfo=tz).timestamp())

    repo.record(
        agent_id="agt-alice",
        user_id=alice_id,
        model="model-a",
        input_tokens=100,
        output_tokens=10,
        ts=day_a,
    )
    repo.record(
        agent_id="agt-alice",
        user_id=alice_id,
        model="model-a",
        input_tokens=200,
        output_tokens=20,
        ts=day_b,
    )
    repo.record(
        agent_id="agt-bob",
        user_id=bob_id,
        model="model-b",
        input_tokens=50,
        output_tokens=5,
        ts=day_b,
    )
    repo.record(
        agent_id="agt-bob",
        user_id=bob_id,
        model="model-b",
        input_tokens=1000,
        output_tokens=100,
        ts=day_c,
    )

    r = await c.get(
        "/api/admin/usage/summary?granularity=total&window=all",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 100 + 10 + 200 + 20 + 50 + 5 + 1000 + 100
    assert r.json()["turns"] == 4

    r = await c.get(
        f"/api/admin/usage/summary?granularity=total&window=all&user_id={alice_id}",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 330
    assert r.json()["turns"] == 2

    r = await c.get(
        f"/api/admin/usage/summary?granularity=total&window=all&user_id={bob_id}",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 1155
    assert r.json()["turns"] == 2

    r = await c.get(
        "/api/admin/usage/summary?granularity=total&window=all&agent_id=agt-bob",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 1155

    r = await c.get(
        f"/api/admin/usage/summary?granularity=total&window=all"
        f"&user_id={bob_id}&agent_id=agt-alice",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 0
    assert r.json()["turns"] == 0

    r = await c.get(
        f"/api/admin/usage/summary?granularity=total&window=day:2026-08-25&user_id={alice_id}",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 220
    assert r.json()["turns"] == 1

    r = await c.get(
        "/api/admin/usage/summary?granularity=total&window=month:2026-09",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 1100
    assert r.json()["turns"] == 1

    r = await c.get(
        "/api/admin/usage/summary?granularity=total&window=range:2026-08-20:2026-08-25",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 100 + 10 + 200 + 20 + 50 + 5
    assert r.json()["turns"] == 3

    r = await c.get(
        f"/api/usage/summary?granularity=total&window=all&as_user={bob_id}",
        headers=alice_auth,
    )
    assert r.status_code == 403

    r = await c.get(
        "/api/usage/summary?granularity=total&window=all",
        headers=bob_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 1155

    r = await c.get(
        f"/api/admin/usage/export.xlsx?window=all&user_id={bob_id}",
        headers={**admin_auth, "Accept-Language": "zh"},
    )
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    detail = wb["明细"]
    headers = [cell.value for cell in detail[1]]
    uid_col = headers.index("用户 ID")
    user_ids = {
        row[uid_col].value
        for row in detail.iter_rows(min_row=2, max_row=detail.max_row - 2)
        if row[uid_col].value is not None
    }
    assert user_ids == {bob_id}


async def test_summary_agent_id_filter_for_user(env: Any) -> None:
    c, srv, _admin_auth, alice_auth, ctx = env
    repo = srv.services.usage_repo
    repo.record(agent_id="agt-a", user_id=ctx["alice_id"], input_tokens=10, output_tokens=5)
    repo.record(agent_id="agt-b", user_id=ctx["alice_id"], input_tokens=100, output_tokens=50)

    r = await c.get(
        "/api/usage/summary?granularity=total&window=all&agent_id=agt-a",
        headers=alice_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 15
    assert r.json()["turns"] == 1

    r = await c.get(
        "/api/usage/summary?granularity=by_agent&window=all&agent_id=agt-b",
        headers=alice_auth,
    )
    assert r.status_code == 200
    buckets = r.json()["buckets"]
    assert len(buckets) == 1
    assert buckets[0]["key"] == "agt-b"
    assert buckets[0]["total_tokens"] == 150


async def test_window_presets_today_yesterday_rolling(env: Any) -> None:
    """Preset windows isolate rows when resolved against a pinned ``now``."""
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from octop.infra.db.repos.usage import resolve_usage_window

    c, srv, admin_auth, _alice_auth, ctx = env
    repo = srv.services.usage_repo
    tz = ZoneInfo("Asia/Shanghai")
    now_ts = int(datetime(2026, 9, 7, 15, 0, tzinfo=tz).timestamp())
    today_ts = int(datetime(2026, 9, 7, 10, 0, tzinfo=tz).timestamp())
    yesterday_ts = int(datetime(2026, 9, 6, 10, 0, tzinfo=tz).timestamp())
    week_ago_ts = int(datetime(2026, 8, 30, 10, 0, tzinfo=tz).timestamp())
    old_ts = int(datetime(2026, 7, 1, 10, 0, tzinfo=tz).timestamp())

    for ts, tokens in (
        (today_ts, 10),
        (yesterday_ts, 20),
        (week_ago_ts, 40),
        (old_ts, 80),
    ):
        repo.record(
            agent_id="y",
            user_id=ctx["alice_id"],
            input_tokens=tokens,
            output_tokens=0,
            ts=ts,
        )

    for window, expected_tokens in (
        ("today", 10),
        ("yesterday", 20),
        # Rolling 7*86400s from 2026-09-07 15:00 → starts 2026-08-31 15:00;
        # the Aug 30 row is outside that window.
        ("last_7d", 10 + 20),
        ("last_30d", 10 + 20 + 40),
        ("all", 10 + 20 + 40 + 80),
    ):
        start, end = resolve_usage_window(window, timezone="Asia/Shanghai", now=now_ts)
        rows = [
            row
            for row in repo.list_detail(
                user_id=ctx["alice_id"],
                window="all",
                timezone="Asia/Shanghai",
            )
            if start <= row.ts < end
        ]
        assert sum(row.total_tokens for row in rows) == expected_tokens, window

    # Calendar windows via admin API (independent of wall-clock ``now``).
    r = await c.get(
        "/api/admin/usage/summary?granularity=total&window=day:2026-09-06"
        f"&user_id={ctx['alice_id']}",
        headers=admin_auth,
    )
    assert r.status_code == 200
    assert r.json()["total_tokens"] == 20


# --- chunk extraction in chat router ---------------------------------------


def test_extract_usage_from_chunk_direct() -> None:
    from octop.api.routers.chat.turn import extract_usage_from_chunk as _extract_usage_from_chunk

    chunk = {"usage": {"input_tokens": 7, "output_tokens": 3}}
    usage = _extract_usage_from_chunk(chunk)
    assert usage is not None
    assert usage["input_tokens"] == 7
    assert usage["uncached_input_tokens"] == 7
    assert usage["cache_read_tokens"] == 0
    assert usage["output_tokens"] == 3


def test_extract_usage_from_state_snapshot_dict_message() -> None:
    from octop.api.routers.chat.turn import extract_usage_from_chunk as _extract_usage_from_chunk

    chunk = {
        "type": "state_snapshot",
        "data": {
            "messages": [
                {"role": "user", "content": "hi"},
                {
                    "role": "assistant",
                    "content": "hello!",
                    "usage_metadata": {"input_tokens": 9, "output_tokens": 4},
                    "response_metadata": {
                        "model_name": "openai:gpt-4o-mini",
                        "token_usage": {
                            "prompt_tokens": 9,
                            "completion_tokens": 4,
                            "prompt_cache_hit_tokens": 6,
                        },
                    },
                },
            ],
        },
    }
    out = _extract_usage_from_chunk(chunk)
    assert out is not None
    assert out["input_tokens"] == 9
    assert out["output_tokens"] == 4
    assert out["cache_read_tokens"] == 6
    assert out["model"] == "openai:gpt-4o-mini"


def test_usage_tracker_accumulates_calls_and_replaces_duplicate_call() -> None:
    from octop.infra.gateway.process.usage_record import UsageTracker

    tracker = UsageTracker()
    tracker.observe(
        {
            "type": "usage",
            "call_id": "call-1",
            "model": "deepseek-v4-pro",
            "usage": {
                "input_tokens": 1_000,
                "uncached_input_tokens": 300,
                "cache_read_tokens": 700,
                "output_tokens": 50,
            },
        }
    )
    tracker.observe(
        {
            "type": "usage",
            "call_id": "call-1",
            "model": "deepseek-v4-pro",
            "usage": {
                "input_tokens": 1_000,
                "uncached_input_tokens": 300,
                "cache_read_tokens": 700,
                "output_tokens": 50,
            },
        }
    )
    tracker.observe(
        {
            "type": "usage",
            "call_id": "call-2",
            "model": "deepseek-v4-pro",
            "usage": {
                "input_tokens": 1_200,
                "uncached_input_tokens": 400,
                "cache_read_tokens": 800,
                "output_tokens": 60,
            },
        }
    )

    assert tracker.usage == {
        "input_tokens": 2_200,
        "uncached_input_tokens": 700,
        "cache_read_tokens": 1_500,
        "cache_write_tokens": 0,
        "output_tokens": 110,
        "reasoning_tokens": 0,
        "total_tokens": 2_310,
        "model": "deepseek-v4-pro",
        "model_calls": 2,
        "last_input_tokens": 1_200,
    }


def test_usage_tracker_explicit_events_supersede_legacy_snapshots() -> None:
    from octop.infra.gateway.process.usage_record import UsageTracker

    tracker = UsageTracker()
    snapshot = {
        "type": "state_snapshot",
        "data": {
            "messages": [
                {
                    "id": "persisted-message-id",
                    "role": "assistant",
                    "usage_metadata": {"input_tokens": 100, "output_tokens": 5},
                }
            ]
        },
    }
    tracker.observe(snapshot)
    tracker.observe(
        {
            "type": "usage",
            "call_id": "stream-call-id",
            "usage": {"input_tokens": 100, "output_tokens": 5},
        }
    )
    tracker.observe(snapshot)

    assert tracker.usage is not None
    assert tracker.usage["input_tokens"] == 100
    assert tracker.usage["model_calls"] == 1


def test_extract_usage_from_state_snapshot_sums_turn_calls() -> None:
    from octop.api.routers.chat.turn import extract_usage_from_chunk as _extract_usage_from_chunk

    chunk = {
        "type": "state_snapshot",
        "data": {
            "messages": [
                {"role": "user", "content": "hi"},
                {
                    "role": "assistant",
                    "usage_metadata": {"input_tokens": 100, "output_tokens": 10},
                },
                {
                    "role": "assistant",
                    "usage_metadata": {"input_tokens": 200, "output_tokens": 20},
                    "response_metadata": {"model_name": "openai:gpt-4o-mini"},
                },
            ],
        },
    }
    out = _extract_usage_from_chunk(chunk)
    assert out is not None
    assert out["input_tokens"] == 300
    assert out["output_tokens"] == 30
    assert out["model"] == "openai:gpt-4o-mini"


def test_extract_usage_returns_none_when_absent() -> None:
    from octop.api.routers.chat.turn import extract_usage_from_chunk as _extract_usage_from_chunk

    assert _extract_usage_from_chunk({"type": "token", "content": "hi"}) is None
    assert _extract_usage_from_chunk({"type": "state_snapshot", "data": {}}) is None
    assert _extract_usage_from_chunk(None) is None  # type: ignore[arg-type]
