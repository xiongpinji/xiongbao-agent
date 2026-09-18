# SPDX-License-Identifier: MIT
"""Office artifact generators (stdlib-first WorkBuddy parity).

Produces Markdown / HTML / CSV always. When optional packages are installed
(``python-docx``, ``openpyxl``, ``python-pptx``), also emits native Office files.
"""

from __future__ import annotations

import csv
import html
import importlib.util
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence


@dataclass
class ArtifactResult:
    kind: str
    path: str
    native: bool
    note: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "kind": self.kind,
            "path": self.path,
            "native": self.native,
            "note": self.note,
        }


def _has(mod: str) -> bool:
    return importlib.util.find_spec(mod) is not None


def write_markdown(path: Path | str, title: str, body: str) -> ArtifactResult:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(f"# {title}\n\n{body.rstrip()}\n", encoding="utf-8")
    return ArtifactResult(kind="markdown", path=str(p), native=True)


def write_html(path: Path | str, title: str, body: str) -> ArtifactResult:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>{html.escape(title)}</title></head><body>"
        f"<h1>{html.escape(title)}</h1>"
        f"<pre>{html.escape(body)}</pre></body></html>\n"
    )
    p.write_text(doc, encoding="utf-8")
    return ArtifactResult(kind="html", path=str(p), native=True)


def write_csv(path: Path | str, rows: Sequence[Sequence[Any]], *, header: Sequence[str] | None = None) -> ArtifactResult:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        if header:
            w.writerow(list(header))
        for row in rows:
            w.writerow(list(row))
    return ArtifactResult(kind="csv", path=str(p), native=True)


def write_xlsx(path: Path | str, rows: Sequence[Sequence[Any]], *, sheet: str = "Sheet1") -> ArtifactResult:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    if _has("openpyxl"):
        from openpyxl import Workbook  # type: ignore

        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for row in rows:
            ws.append(list(row))
        wb.save(p)
        return ArtifactResult(kind="xlsx", path=str(p), native=True)
    # Fallback: CSV sibling
    csv_path = p.with_suffix(".csv")
    write_csv(csv_path, rows)
    p.write_text(
        json.dumps(
            {"fallback": "csv", "csv": str(csv_path), "note": "install openpyxl for native xlsx"},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return ArtifactResult(
        kind="xlsx",
        path=str(csv_path),
        native=False,
        note="openpyxl missing — wrote CSV fallback",
    )


def write_docx(path: Path | str, title: str, paragraphs: Sequence[str]) -> ArtifactResult:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    if _has("docx"):
        from docx import Document  # type: ignore

        doc = Document()
        doc.add_heading(title, level=1)
        for para in paragraphs:
            doc.add_paragraph(para)
        doc.save(p)
        return ArtifactResult(kind="docx", path=str(p), native=True)
    md = p.with_suffix(".md")
    write_markdown(md, title, "\n\n".join(paragraphs))
    return ArtifactResult(
        kind="docx",
        path=str(md),
        native=False,
        note="python-docx missing — wrote Markdown fallback",
    )


def write_pptx(path: Path | str, title: str, slides: Sequence[tuple[str, str]]) -> ArtifactResult:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    if _has("pptx"):
        from pptx import Presentation  # type: ignore

        prs = Presentation()
        # title slide
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        for slide_title, body in slides:
            layout = prs.slide_layouts[1]
            s = prs.slides.add_slide(layout)
            s.shapes.title.text = slide_title
            s.placeholders[1].text = body
        prs.save(p)
        return ArtifactResult(kind="pptx", path=str(p), native=True)
    md = p.with_suffix(".md")
    chunks = [f"# {title}"]
    for st, body in slides:
        chunks.append(f"## {st}\n\n{body}")
    write_markdown(md, title, "\n\n".join(chunks[1:]) if len(chunks) > 1 else "")
    md.write_text("\n\n".join(chunks) + "\n", encoding="utf-8")
    return ArtifactResult(
        kind="pptx",
        path=str(md),
        native=False,
        note="python-pptx missing — wrote Markdown fallback",
    )


def deliver_bundle(
    out_dir: Path | str,
    *,
    title: str,
    body: str,
    table_rows: Sequence[Sequence[Any]] | None = None,
    slides: Sequence[tuple[str, str]] | None = None,
) -> list[ArtifactResult]:
    """Write a full office-like bundle into ``out_dir``."""
    base = Path(out_dir)
    base.mkdir(parents=True, exist_ok=True)
    results = [
        write_markdown(base / "report.md", title, body),
        write_html(base / "report.html", title, body),
        write_docx(base / "report.docx", title, [body]),
    ]
    if table_rows:
        results.append(write_csv(base / "data.csv", table_rows, header=["c1", "c2", "c3"][: len(table_rows[0])]))
        results.append(write_xlsx(base / "data.xlsx", table_rows))
    if slides:
        results.append(write_pptx(base / "deck.pptx", title, slides))
    return results
