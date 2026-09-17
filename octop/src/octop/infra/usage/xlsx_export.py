"""Build localized Excel workbooks for token usage export.

Formatting follows the MiniMax XLSX skill conventions adapted for openpyxl
(runtime API export): bold headers, thousands separators, freeze panes,
auto-filter, TOTAL rows as Excel formulas (blank spacer above so Excel
Sort/Filter does not move them), and readable IO charts.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from octop.i18n import tr
from octop.infra.db.repos.usage import UsageRow
from octop.infra.utils.locale import Locale, normalize_locale

# Match Token Usage dashboard IO colors.
_COLOR_INPUT = "4F6EF7"
_COLOR_OUTPUT = "A06EF7"
_HEADER_FILL = "F3F4F6"
_ZEBRA_FILL = "FAFAFB"
_INT_FORMAT = "#,##0"


def _zoneinfo(timezone: str) -> ZoneInfo:
    try:
        return ZoneInfo(timezone)
    except ZoneInfoNotFoundError:
        return ZoneInfo("UTC")


def _label(locale: Locale, key: str) -> str:
    return tr(f"usage_export.{key}", locale)


def _font(*, bold: bool = False, size: int = 11, color: str | None = None) -> Any:
    from openpyxl.styles import Font

    kwargs: dict[str, Any] = {"bold": bold, "size": size, "name": "Calibri"}
    if color is not None:
        kwargs["color"] = color
    return Font(**kwargs)


def _fill(rgb: str) -> Any:
    from openpyxl.styles import PatternFill

    return PatternFill(fill_type="solid", fgColor=rgb)


def _thin_border() -> Any:
    from openpyxl.styles import Border, Side

    side = Side(style="thin", color="D1D5DB")
    return Border(left=side, right=side, top=side, bottom=side)


def _top_medium_border() -> Any:
    from openpyxl.styles import Border, Side

    thin = Side(style="thin", color="D1D5DB")
    medium = Side(style="medium", color="6B7280")
    return Border(left=thin, right=thin, top=medium, bottom=thin)


def _style_header_row(ws: Any, ncols: int) -> None:
    from openpyxl.styles import Alignment

    header_font = _font(bold=True, size=11)
    header_fill = _fill(_HEADER_FILL)
    border = _thin_border()
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _style_data_cells(
    ws: Any,
    *,
    start_row: int,
    end_row: int,
    ncols: int,
    int_cols: set[int],
) -> None:
    from openpyxl.styles import Alignment

    border = _thin_border()
    zebra = _fill(_ZEBRA_FILL)
    body_font = _font()
    for row in range(start_row, end_row + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row, col)
            cell.font = body_font
            cell.border = border
            if row % 2 == 0:
                cell.fill = zebra
            if col in int_cols and isinstance(cell.value, (int, float)):
                cell.number_format = _INT_FORMAT
                cell.alignment = Alignment(horizontal="right")


def _append_total_row(
    ws: Any,
    *,
    locale: Locale,
    label_col: int,
    sum_cols: list[int],
    data_start: int,
    data_end: int,
    ncols: int,
) -> int:
    """Append a TOTAL row with SUM formulas. Returns the total row index.

    Leaves one blank row between the last data row and TOTAL so Excel's
    Sort / AutoFilter contiguous-region detection does not pull TOTAL
    into the sortable block.
    """
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    # Blank spacer: data_end + 1 stays empty on purpose.
    total_row = data_end + 2
    for col in range(1, ncols + 1):
        cell = ws.cell(total_row, col)
        cell.font = _font(bold=True)
        cell.border = _top_medium_border()
        if col == label_col:
            cell.value = _label(locale, "row_total")
            cell.alignment = Alignment(horizontal="left")
        elif col in sum_cols and data_end >= data_start:
            letter = get_column_letter(col)
            cell.value = f"=SUM({letter}{data_start}:{letter}{data_end})"
            cell.number_format = _INT_FORMAT
            cell.alignment = Alignment(horizontal="right")
    return total_row


def _autosize_columns(ws: Any, ncols: int, *, max_width: int = 36) -> None:
    from openpyxl.utils import get_column_letter

    for col in range(1, ncols + 1):
        width = 10
        for row in ws.iter_rows(min_col=col, max_col=col, max_row=min(ws.max_row, 200)):
            value = row[0].value
            if value is None:
                continue
            width = max(width, min(max_width, len(str(value)) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width


def _format_local_time(ts: int, *, tz: ZoneInfo) -> str:
    """Format unix seconds as ``YYYY-MM-DD HH:mm:ss`` in *tz* (no offset suffix)."""
    return datetime.fromtimestamp(ts, tz=tz).strftime("%Y-%m-%d %H:%M:%S")


def _hide_major_gridlines(axis: Any) -> None:
    """Hide major gridlines (Excel still draws defaults if the element is omitted)."""
    from openpyxl.chart.axis import ChartLines
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    axis.majorGridlines = ChartLines()
    axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(noFill=True))


def _soft_axis_line(axis: Any) -> None:
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="9CA3AF", w=9000))


def _polish_category_axis(axis: Any) -> None:
    axis.axPos = "b"
    axis.delete = False
    axis.majorTickMark = "out"
    axis.minorTickMark = "none"
    axis.tickLblPos = "nextTo"
    axis.majorGridlines = None
    _soft_axis_line(axis)


def _polish_value_axis(axis: Any, *, title: str | None = None) -> None:
    axis.axPos = "l"
    axis.delete = False
    axis.numFmt = _INT_FORMAT
    axis.majorTickMark = "out"
    axis.minorTickMark = "none"
    axis.tickLblPos = "nextTo"
    _hide_major_gridlines(axis)
    _soft_axis_line(axis)
    if title:
        axis.title = title


def _styled_chart_title(text: str) -> Any:
    """Build a larger chart title with a reserved top band.

    Excel chart-title paragraph centering is unreliable; keep a readable
    native title (size/bold) without sheet-cell workarounds.
    """
    from openpyxl.chart.layout import Layout, ManualLayout
    from openpyxl.chart.title import Title
    from openpyxl.drawing.text import (
        CharacterProperties,
        Paragraph,
        ParagraphProperties,
        RegularTextRun,
        RichTextProperties,
    )

    title = Title()
    title.tx.rich.bodyPr = RichTextProperties(anchor="ctr", anchorCtr=True)
    run_props = CharacterProperties(sz=1800, b=True)  # 18pt bold
    para_props = ParagraphProperties(algn="ctr", defRPr=run_props)
    title.tx.rich.paragraphs = [
        Paragraph(pPr=para_props, r=[RegularTextRun(t=text, rPr=run_props)])
    ]
    title.overlay = False
    title.layout = Layout(
        manualLayout=ManualLayout(
            xMode="edge",
            yMode="edge",
            x=0.05,
            y=0.01,
            w=0.9,
            h=0.12,
        )
    )
    return title


def _place_chart_chrome(chart: Any, *, title: str) -> None:
    """Put the title in a reserved top band; keep legend outside the plot."""
    from openpyxl.chart.layout import Layout, ManualLayout

    chart.title = _styled_chart_title(title)
    if chart.legend is not None:
        chart.legend.overlay = False
    chart.layout = Layout(
        manualLayout=ManualLayout(
            xMode="edge",
            yMode="edge",
            x=0.08,
            y=0.16,
            w=0.72,
            h=0.74,
        )
    )


def _add_io_chart(
    ws: Any,
    *,
    title: str,
    categories_col: int,
    data_min_col: int,
    data_max_col: int,
    n_rows: int,
    anchor: str,
    kind: str,
    value_axis_title: str | None = None,
) -> None:
    """Render input/output comparison charts.

    ``kind``:
      - ``line`` — daily trend
      - ``col`` — vertical clustered columns (experts / models)

    Legend sits on the right so category labels stay visible; gridlines are
    suppressed; axes use explicit tick marks and number formats. The title
    sits in a reserved top band (``overlay=False`` + plot layout gap).
    """
    if n_rows < 1:
        return
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.marker import Marker

    data = Reference(
        ws,
        min_col=data_min_col,
        min_row=1,
        max_col=data_max_col,
        max_row=n_rows + 1,
    )
    cats = Reference(ws, min_col=categories_col, min_row=2, max_row=n_rows + 1)
    palette = (_COLOR_INPUT, _COLOR_OUTPUT)

    if kind == "line":
        chart: Any = LineChart()
        chart.legend.position = "r"
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        for idx, series in enumerate(chart.series):
            color = palette[idx % len(palette)]
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = 25000
            series.marker = Marker(symbol="circle", size=5)
            series.marker.graphicalProperties.solidFill = color
            series.marker.graphicalProperties.line.solidFill = color
    else:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.overlap = 0
        chart.gapWidth = 80 if n_rows <= 4 else 140
        chart.legend.position = "r"
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        for idx, series in enumerate(chart.series):
            series.graphicalProperties.solidFill = palette[idx % len(palette)]

    _place_chart_chrome(chart, title=title)
    _polish_category_axis(chart.x_axis)
    _polish_value_axis(chart.y_axis, title=value_axis_title)
    chart.width = 18
    chart.height = max(11, min(15, 9 + n_rows * 0.35))
    ws.add_chart(chart, anchor)


def build_usage_xlsx(
    *,
    rows: list[UsageRow],
    by_day: list[dict[str, Any]],
    by_agent: list[dict[str, Any]],
    by_model: list[dict[str, Any]],
    agent_names: dict[str, str],
    usernames: dict[int, str],
    timezone: str,
    locale: str,
) -> bytes:
    """Build a multi-sheet workbook: detail + by-day/agent/model with charts."""
    from openpyxl import Workbook

    loc = normalize_locale(locale)
    tz = _zoneinfo(timezone)
    wb = Workbook()

    # --- Detail ---
    detail = wb.active
    detail.title = _label(loc, "sheet_detail")
    detail_headers = [
        f"{_label(loc, 'col_time')} ({timezone})",
        _label(loc, "col_user_id"),
        _label(loc, "col_username"),
        _label(loc, "col_agent_id"),
        _label(loc, "col_agent_name"),
        _label(loc, "col_thread_id"),
        _label(loc, "col_model"),
        _label(loc, "col_input_tokens"),
        _label(loc, "col_uncached_input_tokens"),
        _label(loc, "col_cache_read_tokens"),
        _label(loc, "col_cache_write_tokens"),
        _label(loc, "col_output_tokens"),
        _label(loc, "col_reasoning_tokens"),
        _label(loc, "col_total_tokens"),
        _label(loc, "col_model_calls"),
        _label(loc, "col_source"),
    ]
    detail.append(detail_headers)
    detail_int_cols = {8, 9, 10, 11, 12, 13, 14, 15}
    for row in rows:
        detail.append(
            [
                _format_local_time(row.ts, tz=tz),
                row.user_id,
                usernames.get(row.user_id, ""),
                row.agent_id,
                agent_names.get(row.agent_id, row.agent_id),
                row.thread_id,
                row.model,
                row.input_tokens,
                row.uncached_input_tokens,
                row.cache_read_tokens,
                row.cache_write_tokens,
                row.output_tokens,
                row.reasoning_tokens,
                row.total_tokens,
                row.model_calls,
                row.source,
            ]
        )
    ncols = len(detail_headers)
    _style_header_row(detail, ncols)
    if rows:
        _style_data_cells(
            detail,
            start_row=2,
            end_row=1 + len(rows),
            ncols=ncols,
            int_cols=detail_int_cols,
        )
        _append_total_row(
            detail,
            locale=loc,
            label_col=1,
            sum_cols=[8, 9, 10, 11, 12, 13, 14, 15],
            data_start=2,
            data_end=1 + len(rows),
            ncols=ncols,
        )
    from openpyxl.utils import get_column_letter

    last_data_row = max(1, 1 + len(rows))
    detail.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last_data_row}"
    _autosize_columns(detail, ncols)

    def _write_category_sheet(
        *,
        title_key: str,
        chart_title_key: str,
        headers: list[str],
        records: list[list[Any]],
        label_col: int,
        sum_cols: list[int],
        chart_cat_col: int,
        chart_data_min: int,
        chart_data_max: int,
        chart_anchor: str,
        chart_kind: str,
    ) -> None:
        sheet = wb.create_sheet(_label(loc, title_key))
        sheet.append(headers)
        for record in records:
            sheet.append(record)
        n = len(records)
        ncols_local = len(headers)
        int_cols = set(sum_cols)
        _style_header_row(sheet, ncols_local)
        if n:
            _style_data_cells(
                sheet,
                start_row=2,
                end_row=1 + n,
                ncols=ncols_local,
                int_cols=int_cols,
            )
            _append_total_row(
                sheet,
                locale=loc,
                label_col=label_col,
                sum_cols=sum_cols,
                data_start=2,
                data_end=1 + n,
                ncols=ncols_local,
            )
        _autosize_columns(sheet, ncols_local)
        _add_io_chart(
            sheet,
            title=_label(loc, chart_title_key),
            categories_col=chart_cat_col,
            data_min_col=chart_data_min,
            data_max_col=chart_data_max,
            n_rows=n,
            anchor=chart_anchor,
            kind=chart_kind,
            value_axis_title=_label(loc, "axis_tokens"),
        )

    # --- By day (ascending for chart) ---
    day_rows = list(reversed(by_day))
    _write_category_sheet(
        title_key="sheet_by_day",
        chart_title_key="chart_daily",
        headers=[
            _label(loc, "col_date"),
            _label(loc, "col_input_tokens"),
            _label(loc, "col_output_tokens"),
            _label(loc, "col_total_tokens"),
            _label(loc, "col_turns"),
        ],
        records=[
            [
                str(bucket.get("label") or bucket.get("key") or ""),
                int(bucket.get("input_tokens") or 0),
                int(bucket.get("output_tokens") or 0),
                int(bucket.get("total_tokens") or 0),
                int(bucket.get("turns") or 0),
            ]
            for bucket in day_rows
        ],
        label_col=1,
        sum_cols=[2, 3, 4, 5],
        chart_cat_col=1,
        chart_data_min=2,
        chart_data_max=3,
        chart_anchor="G3",
        chart_kind="line",
    )

    # --- By agent ---
    _write_category_sheet(
        title_key="sheet_by_agent",
        chart_title_key="chart_by_agent",
        headers=[
            _label(loc, "col_agent_name"),
            _label(loc, "col_agent_id"),
            _label(loc, "col_input_tokens"),
            _label(loc, "col_output_tokens"),
            _label(loc, "col_total_tokens"),
            _label(loc, "col_turns"),
        ],
        records=[
            [
                agent_names.get(str(bucket.get("key") or ""), str(bucket.get("key") or "")),
                str(bucket.get("key") or ""),
                int(bucket.get("input_tokens") or 0),
                int(bucket.get("output_tokens") or 0),
                int(bucket.get("total_tokens") or 0),
                int(bucket.get("turns") or 0),
            ]
            for bucket in by_agent
        ],
        label_col=1,
        sum_cols=[3, 4, 5, 6],
        chart_cat_col=1,
        chart_data_min=3,
        chart_data_max=4,
        chart_anchor="H3",
        chart_kind="col",
    )

    # --- By model ---
    _write_category_sheet(
        title_key="sheet_by_model",
        chart_title_key="chart_by_model",
        headers=[
            _label(loc, "col_model"),
            _label(loc, "col_input_tokens"),
            _label(loc, "col_output_tokens"),
            _label(loc, "col_total_tokens"),
            _label(loc, "col_turns"),
        ],
        records=[
            [
                str(bucket.get("label") or bucket.get("key") or ""),
                int(bucket.get("input_tokens") or 0),
                int(bucket.get("output_tokens") or 0),
                int(bucket.get("total_tokens") or 0),
                int(bucket.get("turns") or 0),
            ]
            for bucket in by_model
        ],
        label_col=1,
        sum_cols=[2, 3, 4, 5],
        chart_cat_col=1,
        chart_data_min=2,
        chart_data_max=3,
        chart_anchor="G3",
        chart_kind="col",
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
